"""Generate small original synthetic fixtures; never uses organizer data or labels."""
from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from docx import Document
from PIL import Image, ImageDraw, ImageFont

root=Path(__file__).resolve().parents[1]/'public'/'demo'
root.mkdir(parents=True,exist_ok=True)
fields=[('Shipper','NORTHSTAR PAPER LTD'),('Consignee','HARBOR BOOKS LTD'),('Notify Party','HARBOR BOOKS LTD'),('Port of Loading','PORT KLANG'),('Port of Discharge','SINGAPORE'),('Container Count','2'),('Gross Weight','24000 kg')]
si=['SHIPPING INSTRUCTION']+[f'{k}: {v}' for k,v in fields]
c=canvas.Canvas(str(root/'SI-text.pdf'),pagesize=(612,792),invariant=1)
text=c.beginText(50,735);text.setFont('Helvetica',14)
for line in si:text.textLine(line)
c.drawText(text);c.save()
doc=Document();doc.add_paragraph('DRAFT BILL OF LADING')
table=doc.add_table(rows=0,cols=2)
for key,value in fields:
    cells=table.add_row().cells;cells[0].text=key+':';cells[1].text='HARBOR BOOKS INTERNATIONAL LTD' if key=='Consignee' else value
doc.save(root/'BL-table.docx')
# The PDF contains only this raster image, with no text layer.
im=Image.new('RGB',(1600,1500),'white');draw=ImageDraw.Draw(im)
font_path=Path('C:/Windows/Fonts/arial.ttf')
font=ImageFont.truetype(str(font_path) if font_path.exists() else 'DejaVuSans.ttf',42)
for i,line in enumerate(si):draw.text((70,80+i*135),line,fill='black',font=font)
im.save(root/'SI-scan.png')
c=canvas.Canvas(str(root/'SI-scan.pdf'),pagesize=(768,720),invariant=1)
c.drawImage(ImageReader(im),0,0,768,720);c.save()
(root/'README.md').write_text('Original synthetic demonstration documents created for CargoGuard. No organizer records or private answers. SI-text.pdf has a text layer; SI-scan.pdf contains only a raster image; SI-scan.png is the same scan; BL-table.docx holds labeled table cells and one deliberate consignee difference.\n',encoding='utf8')
print('Created original PDF, DOCX, scanned PDF, and PNG demo fixtures.')
