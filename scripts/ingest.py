"""Read-only participant-bundle ingestion. Never opens reference answers.

python scripts/ingest.py --source /path/to/participant-bundle
Optional OCR adapter: set OCR_COMMAND to an executable accepting a file path,
returning UTF-8 text on stdout. No shell interpolation. OCR stays review-required.
"""
import argparse
import hashlib
import json
import os
from pathlib import Path
import shlex
import shutil
import subprocess
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]

def parse_document(path):
    segments = []
    error = None
    method = path.suffix.lower().lstrip('.')
    try:
        if not path.stat().st_size:
            raise ValueError('Document is empty (0 bytes).')
        if method == 'txt':
            segments = [{'location': 'Text document', 'text': path.read_text(encoding='utf-8-sig')}]
        elif method == 'pdf':
            segments = [{'location': f'Page {i+1}', 'text': p.extract_text() or ''} for i,p in enumerate(PdfReader(path).pages)]
        elif method == 'docx':
            doc = Document(path)
            segments = [{'location': 'Document paragraphs', 'text': '\n'.join(p.text for p in doc.paragraphs)}]
            for i, table in enumerate(doc.tables):
                segments.append({'location': f'Table {i+1}', 'text': '\n'.join(' : '.join(c.text for c in row.cells) for row in table.rows)})
        elif method == 'xlsx':
            workbook = load_workbook(path, read_only=True, data_only=True)
            for sheet in workbook:
                rows = []
                for row in sheet.values:
                    values = [str(c) for c in row if c is not None]
                    if values: rows.append(' : '.join(values))
                segments.append({'location': f'Sheet {sheet.title}', 'text': '\n'.join(rows)})
            workbook.close()
        else:
            raise ValueError(f'Unsupported format: {method}')
        if not any(s['text'].strip() for s in segments):
            raise ValueError('No readable text layer; OCR or manual transcription required.')
    except Exception as exc:
        error = str(exc)
    if error and os.environ.get('OCR_COMMAND') and method == 'pdf':
        try:
            result = subprocess.run([*shlex.split(os.environ['OCR_COMMAND']), str(path)], capture_output=True, text=True, encoding='utf8', timeout=60, check=True)
            if result.stdout.strip():
                segments = [{'location': 'OCR output — requires human confirmation', 'text': result.stdout}]
                method = 'ocr'
                error = 'OCR extraction must be confirmed by a reviewer.'
        except Exception as exc:
            error += f' OCR failed: {exc}'
    return {'path': 'attachments/'+path.name, 'name': path.name, 'format': path.suffix[1:].upper(), 'method': method, 'segments': segments, 'text': '\n'.join(s['text'] for s in segments), 'error': error, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}

def ingest(source, output):
    source = source.resolve()
    emails = [json.loads(p.read_text(encoding='utf-8-sig')) for p in sorted((source/'inbox').glob('*.json'))]
    if not emails: raise ValueError('No inbox/*.json records found.')
    docs = {}
    assets = ROOT/'public'/'documents'
    assets.mkdir(parents=True, exist_ok=True)
    for email in emails:
        for name in email.get('attachments', []):
            file = (source/name).resolve()
            if not file.is_relative_to(source/'attachments'): raise ValueError('Attachment path is outside allowed directory.')
            if file.is_file():
                docs[name] = parse_document(file)
                shutil.copyfile(file, assets/file.name)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps({'emails': emails, 'documents': docs, 'provenance': {'source': 'Official participant bundle', 'emails':len(emails), 'documents':len(docs), 'originals_modified':False}}, ensure_ascii=False), encoding='utf8')
    print(json.dumps({'emails':len(emails), 'documents':len(docs), 'unreadable':sum(bool(d['error']) for d in docs.values()), 'output':str(output)}))

if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--source', type=Path, required=True)
    parser.add_argument('--output', type=Path, default=ROOT/'data'/'corpus.json')
    args=parser.parse_args()
    ingest(args.source,args.output)
